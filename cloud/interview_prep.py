"""Builds a fact-checked interview preparation document (Word + PDF).

Usage: python -m cloud.interview_prep <master_profile.json> <prep.json> <output dir>

prep.json: {"code": "PREP-QIDDIYA", "title": "…", "subtitle": "…",
            "sections": [{"heading": "…", "paragraphs": ["…"], "bullets": ["…"], "qa": [{"q": "…", "a": "…"}]}]}
Every paragraph, bullet, question and answer must pass the fact checker.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx.shared import Pt

from app.profile.fact_check import check_text
from app.profile.loader import load_profile
from cloud.app_builder import _bullet, _doc, _heading, to_pdf


class PrepError(ValueError):
    pass


def texts(prep: dict) -> list[str]:
    out = [prep.get("title", ""), prep.get("subtitle", "")]
    for s in prep.get("sections", []):
        out += [s.get("heading", ""), *s.get("paragraphs", []), *s.get("bullets", [])]
        for qa in s.get("qa", []):
            out += [qa["q"], qa["a"]]
    return [t for t in out if t]


def build(profile_path: Path, prep: dict, out: Path, pdf: bool = True) -> list[Path]:
    profile = load_profile(profile_path)
    problems = [f"{v.message} ({v.excerpt})" for t in texts(prep) for v in check_text(t, profile).blocking]
    if problems:
        raise PrepError("; ".join(problems))
    out.mkdir(parents=True, exist_ok=True)
    doc = _doc()
    title = doc.add_paragraph()
    run = title.add_run(prep["title"])
    run.bold, run.font.size = True, Pt(17)
    if prep.get("subtitle"):
        doc.add_paragraph(prep["subtitle"])
    for s in prep.get("sections", []):
        _heading(doc, s["heading"])
        for p in s.get("paragraphs", []):
            doc.add_paragraph(p)
        for b in s.get("bullets", []):
            _bullet(doc, b)
        for qa in s.get("qa", []):
            q = doc.add_paragraph()
            q.paragraph_format.space_before = Pt(6)
            q.add_run(qa["q"]).bold = True
            doc.add_paragraph(qa["a"])
    path = out / f"{prep['code']}_Interview_Prep.docx"
    doc.save(path)
    return [path, to_pdf(path)] if pdf else [path]


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        files = build(Path(argv[0]), json.loads(Path(argv[1]).read_text(encoding="utf-8")), Path(argv[2]))
    except PrepError as exc:
        print(f"BLOCKED: {exc}", file=sys.stderr)
        return 1
    print("\n".join(str(f) for f in files))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
