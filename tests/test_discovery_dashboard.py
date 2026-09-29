import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app

FIXTURE = Path(__file__).parent / "fixtures" / "naukrigulf_alert.html"


def eml() -> bytes:
    return ("From: recommendedjobs@naukrigulf.com\r\nSubject: Top 15 jobs for you today\r\n"
            "Date: Tue, 29 Sep 2026 03:31:26 +0000\r\nMIME-Version: 1.0\r\n"
            "Content-Type: text/html; charset=utf-8\r\n\r\n" + FIXTURE.read_text(encoding="utf-8")).encode("utf-8")


def test_upload_filter_status_and_export(settings):
    with TestClient(create_app(settings)) as client:
        r = client.post("/discovery/upload", files={"file": ("alert.eml", eml(), "message/rfc822")})
        assert r.status_code == 200 and "15 new" in r.text
        page = client.get("/discovery", params={"country": "SA"})
        assert "Elegancia Arabia" in page.text and "Al Futtaim" not in page.text

        vid = client.app.state.vacancies.list(country_code="SA")[0]["id"]
        client.post(f"/discovery/{vid}/status", data={"status": "shortlisted"})
        assert client.app.state.vacancies.get(vid)["status"] == "shortlisted"
        assert "Open job page" in client.get(f"/discovery/{vid}").text

        csv = client.get("/discovery/export/csv").content.decode("utf-8-sig")
        assert csv.splitlines()[0].startswith("Job title,Employer") and len(csv.splitlines()) == 16
        wb = load_workbook(io.BytesIO(client.get("/discovery/export/xlsx", params={"country": "AE"}).content))
        assert wb.active["A1"].value == "Job title" and wb.active.max_row > 2


def test_manual_entry_rejects_login_links_and_detects_duplicates(settings):
    with TestClient(create_app(settings)) as client:
        form = {"title": "Design Manager", "employer": "Example Co", "location": "Riyadh - Saudi Arabia",
                "url": "https://careers.example.com/jobs/42?utm=x", "experience": "10 - 15 Years"}
        assert "Vacancy added" in client.post("/discovery/manual", data=form).text
        assert "already in the list" in client.post("/discovery/manual", data=form).text
        bad = dict(form, url="https://example.com/autologin?token=abc")
        assert "not saved" in client.post("/discovery/manual", data=bad).text
        item = client.app.state.vacancies.list()[0]
        assert item["apply_url"] == "https://careers.example.com/jobs/42"


def test_gmail_import_without_setup_explains(settings):
    with TestClient(create_app(settings)) as client:
        r = client.post("/discovery/import-gmail", data={"days": 14})
        assert "Gmail is not set up" in r.text


def test_formula_injection_is_neutralised(settings):
    with TestClient(create_app(settings)) as client:
        client.post("/discovery/manual", data={"title": "=HYPERLINK(\"http://evil\")", "employer": "X"})
        csv = client.get("/discovery/export/csv").content.decode("utf-8-sig")
        assert "'=HYPERLINK" in csv


def test_boards_page(settings):
    with TestClient(create_app(settings)) as client:
        r = client.post("/discovery/boards/add", data={"employer": "Example", "ats": "lever", "token": "example"})
        assert "added to the watch-list" in r.text
        assert "Example" in client.get("/discovery/boards").text
