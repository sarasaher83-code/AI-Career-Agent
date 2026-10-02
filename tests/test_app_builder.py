from pathlib import Path

import pytest
from docx import Document

from cloud.app_builder import TailoringError, build, validate

FIXTURE = Path(__file__).parent / "fixtures" / "sample_profile.json"


def tailoring(**overrides):
    t = {
        "code": "J001", "job_title": "Senior Design Manager", "employer": "Example Co",
        "headline": "Head of Design Management",
        "summary": "Design management leader responsible for 12 residential projects totalling 90,000+ sqm in Riyadh.",
        "competencies": [], "highlights": {"0": [1, 0]}, "projects": [0],
        "letter": {"paragraphs": ["I founded the design management function and led 12 residential projects."]},
    }
    t.update(overrides)
    return t


def test_builds_cv_and_letter_with_verbatim_highlights(tmp_path, profile):
    files = build(FIXTURE, tailoring(), tmp_path, pdf=False)
    assert [f.suffix for f in files] == [".docx", ".docx"]
    text = "\n".join(p.text for p in Document(files[0]).paragraphs)
    assert "TEST CANDIDATE" in text and "Jan 2022 – Present" in text
    first, second = profile.raw["experience"][0]["highlights"][1], profile.raw["experience"][0]["highlights"][0]
    assert text.index(first) < text.index(second)                 # chosen order kept
    assert "Harbour Tower" in text and "7777" not in text         # private preferences never printed
    letter = "\n".join(p.text for p in Document(files[1]).paragraphs)
    assert "Re: Senior Design Manager, Example Co" in letter
    assert "candidate.example.org" in text and "Portfolio: https://candidate.example.org" in letter


def test_invented_claims_are_blocked(profile):
    problems = validate(profile, tailoring(summary="Managed a SAR 2 billion portfolio across 40 projects as an MBA."))
    joined = " ".join(problems)
    assert "Currency" in joined and "40" in joined and "MBA" in joined


def test_indices_and_competencies_must_exist(profile):
    problems = validate(profile, tailoring(highlights={"0": [9]}, projects=[5], competencies=["Rocket Science"]))
    assert len(problems) == 3


def test_build_refuses_blocked_tailoring(tmp_path):
    with pytest.raises(TailoringError):
        build(FIXTURE, tailoring(summary="Leader with 40 years of experience."), tmp_path, pdf=False)
