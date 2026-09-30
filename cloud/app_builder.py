"""Builds a tailored, ATS-friendly CV and cover letter (Word + PDF) for one job.

The tailoring file chooses and orders content; it cannot add experience:
  * achievements are picked by index from the master profile (text is copied verbatim),
  * the summary and cover letter are free text, and must pass the fact checker.

Usage: python -m cloud.app_builder <master_profile.json> <tailoring.json> <output dir>

tailoring.json:
{
  "code": "J001", "job_title": "Senior Design Manager", "employer": "Qiddiya",
  "headline": "Head of Design Management Operations & Technical Office | Real Estate Development",
  "summary": "…", "competencies": ["…"],           # competencies must exist in profile skills.core
  "highlights": {"0": [0, 3, 7], "1": [0, 2]},      # role index -> highlight indices, in order
  "projects": [0, 2, 3],                            # project indices, in order
  "letter": {"salutation": "Dear Hiring Team,", "paragraphs": ["…", "…"], "closing": "Kind regards,"}
}
"""

from __future__ import annotations

import html
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.profile.fact_check import check_text
from app.profile.loader import load_profile

INK = RGBColor(0x14, 0x14, 0x14)
GOLD = RGBColor(0x8C, 0x71, 0x40)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class TailoringError(ValueError):
    pass


def _period(start: str | None, end: str | None) -> str:
    def fmt(value):
        if not value:
            return "Present"
        year, _, month = str(value).partition("-")
        return f"{MONTHS[int(month) - 1]} {year}" if month else year
    return f"{fmt(start)} – {fmt(end)}"


def _doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(10.5), INK
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(46)
    return doc


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
    run = p.add_run(text.upper())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10.5), GOLD


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)


def _contact_block(doc: Document, identity: dict, headline: str) -> None:
    name = doc.add_paragraph()
    run = name.add_run(identity["name"].upper())
    run.bold, run.font.size = True, Pt(18)
    title = doc.add_paragraph()
    title.add_run(headline).font.size = Pt(11)
    contact = doc.add_paragraph()
    parts = [identity.get("base", ""), identity.get("phone", ""), identity.get("email", ""),
             identity.get("linkedin", "").replace("https://www.", "")]
    contact.add_run(" | ".join(p for p in parts if p)).font.size = Pt(9.5)


def validate(profile, tailoring: dict) -> list[str]:
    """Returns every problem; an empty list means the tailoring may be built."""
    raw, problems = profile.raw, []
    roles = raw["experience"]
    for role_key, picks in tailoring.get("highlights", {}).items():
        index = int(role_key)
        if index >= len(roles):
            problems.append(f"role {index} does not exist")
            continue
        for pick in picks:
            if not 0 <= pick < len(roles[index].get("highlights", [])):
                problems.append(f"role {index} has no highlight {pick}")
    for pick in tailoring.get("projects", []):
        if not 0 <= pick < len(raw.get("projects", [])):
            problems.append(f"project {pick} does not exist")
    core = set(raw.get("skills", {}).get("core", []))
    for item in tailoring.get("competencies", []):
        if item not in core:
            problems.append(f'competency "{item}" is not in the profile')
    letter = tailoring.get("letter", {})
    free_text = "\n".join([tailoring.get("summary", ""), tailoring.get("headline", ""),
                           *letter.get("paragraphs", [])])
    for v in check_text(free_text, profile).blocking:
        problems.append(f"fact check: {v.message} ({v.excerpt})")
    return problems


def build_cv(profile, tailoring: dict, out: Path) -> Path:
    raw = profile.raw
    doc = _doc()
    _contact_block(doc, raw["identity"], tailoring["headline"])

    _heading(doc, "Professional Summary")
    doc.add_paragraph(tailoring["summary"])

    if tailoring.get("competencies"):
        _heading(doc, "Core Competencies")
        doc.add_paragraph(" • ".join(tailoring["competencies"]))

    _heading(doc, "Professional Experience")
    for index, role in enumerate(raw["experience"]):
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(6)
        head.add_run(role["title"]).bold = True
        head.add_run(f"  |  {_period(role.get('start'), role.get('end'))}")
        sub = doc.add_paragraph()
        details = [role["employer"], role.get("location", "")]
        if role.get("reports_to"):
            details.append(f"Reporting to the {role['reports_to']}")
        sub.add_run(" | ".join(d for d in details if d)).italic = True
        picks = tailoring.get("highlights", {}).get(str(index))
        chosen = [role["highlights"][i] for i in picks] if picks is not None else role.get("highlights", [])
        for text in chosen:
            _bullet(doc, text)

    projects = [raw["projects"][i] for i in tailoring.get("projects", [])]
    if projects:
        _heading(doc, "Selected Projects")
        for p in projects:
            status = f" ({p['status']})" if p.get("status") else ""
            _bullet(doc, f"{p['name']}, {p['location']}{status}: {p['desc']}")

    _heading(doc, "Education")
    for e in raw["education"]:
        _bullet(doc, f"{e['degree']}, {e['institution']} ({e['start']}–{e['end']})")

    _heading(doc, "Certifications & Training")
    for c in raw.get("certifications_and_training", []):
        date = f" ({c['date']})" if c.get("date") else ""
        _bullet(doc, f"{c['name']}, {c['issuer']}{date}")

    _heading(doc, "Technical & AI Skills")
    skills = raw.get("skills", {})
    doc.add_paragraph("; ".join(skills.get("technical", [])))
    doc.add_paragraph("AI: " + ", ".join(skills.get("ai", [])[:6]))

    _heading(doc, "Languages")
    doc.add_paragraph(" | ".join(f"{l['language']}: {l['level']}" for l in raw["identity"].get("languages", [])))

    path = out / f"{tailoring['code']}_CV_{raw['identity']['name'].replace(' ', '_')}.docx"
    doc.save(path)
    return path


def build_letter(profile, tailoring: dict, out: Path) -> Path:
    raw, letter = profile.raw, tailoring["letter"]
    doc = _doc()
    _contact_block(doc, raw["identity"], tailoring["headline"])
    doc.add_paragraph()
    doc.add_paragraph(f"Re: {tailoring['job_title']}, {tailoring['employer']}").runs[0].bold = True
    doc.add_paragraph(letter.get("salutation", "Dear Hiring Team,"))
    for paragraph in letter["paragraphs"]:
        p = doc.add_paragraph(paragraph)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(8)
    doc.add_paragraph(letter.get("closing", "Kind regards,"))
    doc.add_paragraph(raw["identity"]["name"])
    path = out / f"{tailoring['code']}_Cover_Letter_{raw['identity']['name'].replace(' ', '_')}.docx"
    doc.save(path)
    return path


def _docx_as_html(docx_path: Path) -> str:
    """The generated document as simple HTML, paragraph by paragraph, keeping bold, italics, bullets and headings."""
    parts, in_list = [], False
    for p in Document(docx_path).paragraphs:
        runs = []
        for r in p.runs:
            text = html.escape(r.text)
            if r.bold:
                text = f"<b>{text}</b>"
            if r.italic:
                text = f"<i>{text}</i>"
            runs.append(text)
        body = "".join(runs)
        is_bullet = p.style.name.startswith("List Bullet")
        if in_list and not is_bullet:
            parts.append("</ul>")
            in_list = False
        if is_bullet:
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{body}</li>")
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=10.5)
        colour = next((r.font.color.rgb for r in p.runs if r.font.color and r.font.color.type is not None), None)
        if size >= 16:
            parts.append(f"<h1>{body}</h1>")
        elif colour == GOLD:
            parts.append(f"<h2>{body}</h2>")
        else:
            parts.append(f"<p>{body or '&nbsp;'}</p>")
    if in_list:
        parts.append("</ul>")
    return "".join(parts)


PDF_CSS = """
body { font-family: sans-serif; font-size: 10pt; color: #141414; }
h1 { font-size: 17pt; margin: 0 0 2pt 0; }
h2 { font-size: 10pt; color: #8C7140; margin: 9pt 0 2pt 0; }
p { margin: 0 0 3pt 0; }
ul { margin: 0 0 3pt 0; }
li { margin: 0 0 1pt 0; }
"""


def to_pdf(docx_path: Path) -> Path:
    """Render the document to an A4 PDF with PyMuPDF (no office suite needed)."""
    import pymupdf

    out = docx_path.with_suffix(".pdf")
    story = pymupdf.Story(html=_docx_as_html(docx_path), user_css=PDF_CSS)
    writer = pymupdf.DocumentWriter(str(out))
    page, margin = pymupdf.paper_rect("a4"), 50
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(page + (margin, margin, -margin, -margin))
        story.draw(device)
        writer.end_page()
    writer.close()
    compact = pymupdf.open(out)                  # embed only the glyphs used: smaller files for uploads and email
    compact.subset_fonts()
    compact.save(out.with_suffix(".tmp"), garbage=4, deflate=True)
    compact.close()
    out.with_suffix(".tmp").replace(out)
    if not out.exists() or out.stat().st_size == 0:
        raise RuntimeError(f"PDF was not created for {docx_path.name}")
    return out


def build(profile_path: Path, tailoring: dict, out: Path, pdf: bool = True) -> list[Path]:
    profile = load_profile(profile_path)
    problems = validate(profile, tailoring)
    if problems:
        raise TailoringError("; ".join(problems))
    out.mkdir(parents=True, exist_ok=True)
    files = [build_cv(profile, tailoring, out), build_letter(profile, tailoring, out)]
    if pdf:
        files += [to_pdf(f) for f in files]
    return files


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        files = build(Path(argv[0]), json.loads(Path(argv[1]).read_text(encoding="utf-8")), Path(argv[2]))
    except TailoringError as exc:
        print(f"BLOCKED: {exc}", file=sys.stderr)
        return 1
    print("\n".join(str(f) for f in files))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
