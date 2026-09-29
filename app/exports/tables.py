"""CSV and Excel exports for any dashboard table."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="EFE6D3")   # soft matte gold
HEADER_FONT = Font(bold=True, color="141414")


def _safe(value):
    """Neutralise spreadsheet formulas in text that came from outside (e.g. a job title starting with '=')."""
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(rows: list[dict], columns: list[tuple[str, str]]) -> bytes:
    """UTF-8 with BOM so Excel on Windows shows Arabic text correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([_safe(row.get(key)) for key, _ in columns])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def to_xlsx(rows: list[dict], columns: list[tuple[str, str]], sheet_title: str = "Export") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([_safe(row.get(key)) for key, _ in columns])
    for index, (key, label) in enumerate(columns, start=1):
        width = max([len(str(label))] + [len(str(r.get(key) or "")) for r in rows[:200]])
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
